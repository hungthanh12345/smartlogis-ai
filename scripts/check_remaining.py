import docx
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

file_path = r"C:\Users\dat17\Downloads\BAO_CAO_DU_AN_QUAN_LY_KHO  (4).docx"
doc = docx.Document(file_path)

print("--- REMAINING 'KẾ TOÁN' OR '3 VAI TRÒ' OR '3 ACTORS' ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text
    if any(k in txt for k in ['3 vai trò', '3 Actors', '3 actors', 'Kế toán']):
        print(f"P[{i}]: {txt.strip()}")

for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        row_txt = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
        if any(k in row_txt for k in ['3 vai trò', '3 Actors', '3 actors']):
            print(f"T[{t_idx}] R[{r_idx}]: {row_txt[:120]}")
