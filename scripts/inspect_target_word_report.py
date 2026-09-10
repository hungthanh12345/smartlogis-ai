import sys
import os
import zipfile
import docx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

file_path = r"C:\Users\dat17\Downloads\BAO_CAO_DU_AN_QUAN_LY_KHO  (4).docx"
print(f"Checking file: {file_path}")
if not os.path.exists(file_path):
    print("File does not exist!")
    sys.exit(1)

with zipfile.ZipFile(file_path, 'r') as z:
    media_files = [n for n in z.namelist() if n.startswith('word/media/')]
    print(f"Total media files: {len(media_files)}")
    for mf in sorted(media_files):
        info = z.getinfo(mf)
        print(f"  {mf}: {info.file_size} bytes")

doc = docx.Document(file_path)
print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- Key Paragraphs mentioning Roles / Diagrams ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if any(k in txt.lower() for k in ['3 vai trò', '3 actors', 'kế toán', 'thủ kho', 'hình 2.', 'hình 4.']):
        if len(txt) > 0:
            print(f"P[{i}]: {txt[:120]}")

print("\n--- Key Table Rows mentioning Roles ---")
for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        row_txt = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
        if any(k in row_txt.lower() for k in ['kế toán', 'thủ kho', '3 vai trò', 'ketoan', 'thukho']):
            print(f"T[{t_idx}] R[{r_idx}]: {row_txt[:120]}")
