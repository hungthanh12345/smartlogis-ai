import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r"C:\Users\dat17\OneDrive\Documents\BAO_CAO_DU_AN_QUAN_LY_KHO  (4).docx"
doc = docx.Document(doc_path)
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

print("\n--- MATCHING PARAGRAPHS ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text
    if any(k in txt.lower() for k in ['kế toán', 'thủ kho', '3 vai trò', 'vai trò', 'thukho', 'ketoan']):
        if len(txt.strip()) > 0:
            print(f"P[{i}]: {txt.strip()}")

print("\n--- MATCHING TABLES ---")
for t_idx, tbl in enumerate(doc.tables):
    found = False
    for r_idx, row in enumerate(tbl.rows):
        row_txt = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
        if any(k in row_txt.lower() for k in ['kế toán', 'thủ kho', 'vai trò', 'thukho', 'ketoan']):
            if not found:
                print(f"\nTable {t_idx} (rows: {len(tbl.rows)}, cols: {len(tbl.columns)}):")
                found = True
            print(f"  Row {r_idx}: {row_txt[:140]}")
