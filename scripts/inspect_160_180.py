import docx
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

file_path = r"C:\Users\dat17\Downloads\BAO_CAO_DU_AN_QUAN_LY_KHO  (4).docx"
doc = docx.Document(file_path)

for i in range(160, 180):
    if i < len(doc.paragraphs):
        t = doc.paragraphs[i].text.strip()
        if t:
            print(f"P[{i}]: {t}")
